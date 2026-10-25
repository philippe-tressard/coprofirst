"""Lecture d'un classeur Excel d'import — la mécanique, écrite une fois.

Les trois imports (lots, télécommandes, vigiks) partageaient **trois fonctions
identiques au mot près** : `normaliser`, `importer_depuis_bytes` et
`importer_depuis_fichier`. Seuls le nom dans la docstring et le traitement de
lignes appelé changeaient. Cinquante-sept lignes recopiées trois fois, dont la
gestion d'`openpyxl` absent, l'ouverture en lecture seule, la fermeture du
classeur et la transaction.

C'est la duplication décrite par `standards/02-factorisation.md` §2 : la mécanique
d'un import n'a aucune raison de dépendre de ce qu'on importe. Ce qui diffère
vraiment — comment interpréter les lignes — reste dans chaque module, sous la
forme d'un `traiter(rows, session, remplacer)` passé en paramètre.

Ce qui n'est **pas** ici, volontairement : la lecture des colonnes, la résolution
des bâtiments, la détection des doublons. Ce sont trois règles métier distinctes
qui se ressemblent peu ; les fondre créerait le couplage que le §4 du même
standard met en garde.
"""

import io
import unicodedata
from pathlib import Path
from typing import Callable, Optional

from sqlmodel import Session, select

from app import contexte

#: Les porteurs d'accès qui ne sont PAS des résidents — parkings publics, accès
#: de la mairie, prestataire d'entretien. Une ligne à leur nom n'a personne à
#: rapprocher : l'import la met de côté au lieu de la laisser en attente
#: éternelle dans l'écran d'administration.
#:
#: 🔴 Ces trois noms étaient écrits DEUX fois (14/09/2026, #779) — une fois dans
#: chaque importeur — et le classeur des télécommandes en ajoute deux qui lui
#: sont propres. C'est un NOYAU commun enrichi localement, la forme même que la
#: recopie détruit : ajouter un quatrième non-résident demandait deux gestes, et
#: le second s'oublie.
#:
#: ⚠️ Ces valeurs sont des noms propres de CETTE copropriété, et elles n'ont
#: rien à faire dans le code — même famille que « 5Hostachy » (v1.36.7) et que
#: « 4 SAISONS » (#953). Les réunir ici est le PRÉALABLE à les administrer : on
#: ne déplace pas vers un écran ce qui est écrit à deux endroits.
NOMS_NON_RESIDENTS = frozenset(
    {
        "PARKINGS PUBLIQUES",
        "0. ACCES MAIRIE",
        "ATPE",
    }
)

#: Signature du traitement propre à chaque import : (lignes, session, remplacer) → stats.
Traitement = Callable[[list, Session, bool], dict]

_ERREUR_OPENPYXL = "openpyxl n'est pas installé. Ajouter au requirements.txt."


def normaliser(s: Optional[str]) -> str:
    """Normalise une chaîne : majuscules, sans accents, espaces normalisés.

    Sert à comparer des noms saisis à la main dans trois fichiers différents —
    « Dupont-Martin », « DUPONT MARTIN » et « Dupont  Martin » doivent
    s'apparier. Écrite trois fois à l'identique jusqu'au 08/08/2026.
    """
    if not s:
        return ""
    s = s.strip().upper()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return " ".join(s.split())


def _lignes(source) -> list:
    """Toutes les lignes de la première feuille, classeur refermé derrière.

    `read_only` + `data_only` : on lit des valeurs, jamais des formules, et sans
    charger le classeur entier en mémoire — un import de lots fait plusieurs
    milliers de lignes.
    """
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise RuntimeError(_ERREUR_OPENPYXL)

    classeur = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        return list(classeur.active.iter_rows(values_only=True))
    finally:
        #  Fermeture dans un `finally` : en `read_only`, openpyxl garde un
        #  descripteur ouvert sur l'archive. Les trois copies d'origine
        #  fermaient après la lecture, donc jamais si celle-ci levait.
        classeur.close()


def lignes_du_classeur(contenu: bytes) -> list:
    """Les lignes d'un classeur reçu en octets — pour un import qui n'a pas la forme
    « une ligne, un objet en attente » de `importeurs` (l'état descriptif, #2213)."""
    return _lignes(io.BytesIO(contenu))


def feuille_du_classeur(contenu: bytes, nom: str) -> list | None:
    """Les lignes de la feuille `nom` (casse et accents ignorés), `None` si elle manque —
    la feuille « Clés » facultative de l'état descriptif (#2251)."""
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise RuntimeError(_ERREUR_OPENPYXL)

    classeur = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    try:
        for feuille in classeur.worksheets:
            if normaliser(feuille.title) == normaliser(nom):
                return list(feuille.iter_rows(values_only=True))
        return None
    finally:
        classeur.close()


def importer_bytes(
    contenu: bytes,
    session: Session,
    remplacer: bool,
    traiter: Traitement,
) -> dict:
    """Import depuis des octets en mémoire — chemin du téléversement HTTP.

    La session est celle de la requête : c'est l'appelant qui décide de sa
    portée, on se contente de valider la transaction.
    """
    stats = traiter(_lignes(io.BytesIO(contenu)), session, remplacer)
    session.commit()
    return stats


def importer_fichier(chemin: str, remplacer: bool, traiter: Traitement) -> dict:
    """Import depuis un fichier sur disque — chemin du script en ligne de commande.

    Ouvre sa **propre** session : appelé hors requête, il n'y en a aucune à
    reprendre.
    """
    path = Path(chemin)
    if not path.exists():
        raise FileNotFoundError(f"Fichier introuvable : {chemin}")

    rows = _lignes(str(path))
    with contexte.nouvelle_session() as session:
        stats = traiter(rows, session, remplacer)
        session.commit()
    return stats


def importeurs(traiter: Traitement) -> tuple[Callable[..., dict], Callable[..., dict]]:
    """Les deux points d'entrée d'un import — `(depuis_bytes, depuis_fichier)` (#1564).

    Chacun des trois modules d'import redéfinissait ces deux fonctions, à la
    docstring près : seul le `traiter` passé à `importer_bytes` / `importer_fichier`
    changeait. Le module n'écrit plus que ce traitement, et nomme ses deux
    importeurs en une ligne, juste après lui.
    """

    def depuis_bytes(contenu: bytes, session: Session, remplacer: bool = False) -> dict:
        """Import depuis des octets en mémoire (téléversement HTTP)."""
        return importer_bytes(contenu, session, remplacer, traiter)

    def depuis_fichier(chemin: str, remplacer: bool = False) -> dict:
        """Importe depuis un xlsx (script en ligne de commande).

        Rend un dict aux clés ``importes``, ``ignores``, ``doublons``, ``erreurs``.
        """
        return importer_fichier(chemin, remplacer, traiter)

    return depuis_bytes, depuis_fichier


def purger_staging(session: Session, modele, statuts) -> int:
    """Vide le staging avant un réimport « Remplacer » — et rend ce qu'on efface.

    🔴 Ce geste était recopié **trois fois** (lots, télécommandes, vigiks), et les
    trois copies ne purgeaient **pas la même chose** sans qu'un seul commentaire
    ne dise si la divergence était voulue (#824) :

    ==============  ==========================  ====================================
    Import          Prédicat d'origine          Ce qui disparaissait
    ==============  ==========================  ====================================
    télécommandes   ``statut == en_attente``    le non-traité
    vigiks          ``statut == en_attente``    le non-traité
    **lots**        ``statut != resolu``        le non-traité **et le mis de côté**
    ==============  ==========================  ====================================

    La troisième ligne est un défaut. ``ignore`` veut dire *« l'admin a choisi
    d'écarter cette ligne »* : réimporter le même fichier avec « Remplacer » la
    **ressuscitait**, sans rien dire — la forme exacte du seed des périmètres qui
    annulait les suppressions de l'administration (13/08/2026). Une décision
    d'administration n'est pas une donnée d'import, et ne se réécrit pas depuis
    un classeur.

    ⚠️ Le paramètre nomme ce qu'on **efface**, jamais ce qu'on garde. Un statut
    ajouté à l'énumération est alors préservé par défaut : l'oubli va vers la
    conservation, pas vers la perte. C'est le seul sens dans lequel un oubli est
    rattrapable.

    Chaque appelant déclare ses statuts **avec son motif** — la divergence qui
    reste après #824 est celle des lots, qui purgent aussi les rapprochements à
    demi faits, et elle est écrite là où elle se décide.
    """
    condamnes = session.exec(select(modele).where(modele.statut.in_(list(statuts)))).all()
    for ligne in condamnes:
        session.delete(ligne)
    #  `flush` AVANT les insertions : sans lui, SQLAlchemy peut ordonner les
    #  suppressions après elles et effacer ce qu'on vient d'écrire.
    session.flush()
    return len(condamnes)


#  ── Le vocabulaire de la colonne « TYPE » d'un classeur de lots ──────────────
#
#  🔴 Ces trois tables et les deux convertisseurs qui les lisent étaient écrits
#  DEUX fois (#829) : au niveau module dans `routers/lots.py`, et **dans un corps
#  de fonction** dans `utils/auto_match_service.py`. La seconde copie existait
#  parce que le `_norm` du même fichier faisait autre chose que celui de
#  `lots.py` — il avait fallu réécrire la bonne version sous le nom `_norm2`
#  pour l'avoir sous la main.
#
#  Un nom qui ment produit une copie, pas une erreur. C'est pourquoi le
#  normaliseur de comparaison de noms s'appelle désormais `_cle_de_nom`.
#
#  ⚠️ Ce vocabulaire vit ICI et non dans un routeur : il décrit ce qu'un
#  CLASSEUR contient, comme `normaliser` juste au-dessus. Un routeur qui le
#  porte le rend invisible à tout autre lecteur du même classeur — ce qui est
#  exactement ce qui s'est produit.

#: Les préfixes de la colonne TYPE qui désignent un emplacement, pas un logement.
TYPE_PARKING = {"PS"}
TYPE_CAVE = {"CA"}

#: Les mentions d'étage du classeur, en entiers. Les sous-sols sont négatifs.
ETAGE_PAR_MENTION: dict[str, int] = {
    "RDC": 0,
    "1ER": 1,
    "1SS": -1,
    "2EME": 2,
    "2SS": -2,
    "3EME": 3,
    "3SS": -3,
    "4EME": 4,
    "5EME": 5,
}

#: Les valeurs de TYPE qui ne qualifient PAS un appartement — « AP » est le cas
#: général, « DIV » et « LC » des mentions administratives sans type de logement.
_TYPE_SANS_QUALIFICATION = ("AP", "DIV", "LC", "")


def type_de_lot(type_raw: Optional[str]):
    """La colonne TYPE d'un classeur → `(TypeLot, type d'appartement | None)`.

    L'import est différé : `app.models.core` importe la base, et ce module est
    chargé par des scripts en ligne de commande qui n'en ont pas besoin.
    """
    from app.models.core import TypeLot

    t = normaliser(type_raw)
    if t in TYPE_PARKING:
        return TypeLot.parking, None
    if t in TYPE_CAVE:
        return TypeLot.cave, None
    return TypeLot.appartement, (t if t not in _TYPE_SANS_QUALIFICATION else None)


def etage_de_lot(etage_raw: Optional[str]) -> Optional[int]:
    """La colonne ÉTAGE d'un classeur → un entier, ou `None` si illisible.

    ⚠️ Une mention inconnue rend `None` et **n'échoue pas** : un classeur qui
    invente une mention doit importer le lot sans étage, pas bloquer la ligne.
    """
    return ETAGE_PAR_MENTION.get(normaliser(etage_raw)) if etage_raw else None
